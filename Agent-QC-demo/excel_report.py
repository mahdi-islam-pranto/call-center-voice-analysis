"""
Builds the downloadable Excel QC report from the /agent-performance response JSON.

Sheets:
  Overview              - agent-level score, summary, strengths, weaknesses, metric averages
  Scores by Recording   - one row per file, one column per performance type (+ total, %, call summary)
  Detailed Justification - one row per (file, performance type) with score and justification
"""
from datetime import datetime
from io import BytesIO
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SCORE_MAX = 10

HEADER_FILL = PatternFill("solid", start_color="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LABEL_FONT = Font(bold=True)
TITLE_FONT = Font(bold=True, size=16)
SECTION_FONT = Font(bold=True, size=12, color="1F3864")
AVERAGE_FILL = PatternFill("solid", start_color="D9E1F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top")


def _set(ws, row: int, col: int, value: Any, formula: bool = False, **style):
    """Writes a cell. Unless formula=True, text starting with '=' is kept as plain text so LLM output / file names never run as formulas."""
    cell = ws.cell(row=row, column=col, value=value)
    if not formula and isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"
    for key, val in style.items():
        setattr(cell, key, val)
    return cell


def _header_row(ws, row: int, headers: list[str]):
    for col, text in enumerate(headers, start=1):
        _set(ws, row, col, text, font=HEADER_FONT, fill=HEADER_FILL, alignment=WRAP, border=BORDER)


def _score_color_scale(ws, cell_range: str):
    ws.conditional_formatting.add(
        cell_range,
        ColorScaleRule(
            start_type="num", start_value=0, start_color="F8696B",
            mid_type="num", mid_value=SCORE_MAX / 2, mid_color="FFEB84",
            end_type="num", end_value=SCORE_MAX, end_color="63BE7B",
        ),
    )


def _bullets(items: Optional[list[str]]) -> str:
    return "\n".join(f"• {i}" for i in items or []) or "-"


def _calls_from_report(report: dict) -> list[dict]:
    """Per-call results; older API versions only return the combined report, so treat it as one row."""
    calls = report.get("call_evaluations") or []
    if calls:
        return calls
    return [
        {
            "file_name": "All recordings (combined)",
            "metrics": report.get("metrics") or [],
            "call_summary": report.get("performance_summary", ""),
        }
    ]


def _overview_sheet(ws, report: dict, response: dict, calls: list[dict], generated_at: datetime):
    ws.title = "Overview"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 90

    _set(ws, 1, 1, "Agent Performance QC Report", font=TITLE_FONT)
    failed = response.get("failed_urls") or report.get("failed_urls") or []
    info = [
        ("Generated at", generated_at.strftime("%Y-%m-%d %H:%M")),
        ("Agent name", report.get("agent_name") or "Not identified"),
        ("Recordings evaluated", len(report.get("call_evaluations") or []) or "-"),
        ("Overall score", f"{report.get('total_score', 0)} / {report.get('max_possible_score', 0)}"),
        ("Overall percentage", (report.get("percentage") or 0) / 100),
        ("Failed recordings", "\n".join(failed) or "None"),
    ]
    row = 3
    for label, value in info:
        _set(ws, row, 1, label, font=LABEL_FONT)
        cell = _set(ws, row, 2, value, alignment=Alignment(horizontal="left", vertical="top", wrap_text=True))
        if label == "Overall percentage":
            cell.number_format = "0.0%"
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        row += 1

    row += 1
    for label, value in [
        ("Performance summary", report.get("performance_summary") or "-"),
        ("Strengths", _bullets(report.get("strengths"))),
        ("Weaknesses", _bullets(report.get("weaknesses"))),
        ("Improvement suggestions", _bullets(report.get("improvement_suggestions"))),
    ]:
        _set(ws, row, 1, label, font=LABEL_FONT, alignment=WRAP)
        _set(ws, row, 2, value, alignment=WRAP)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        lines = sum(1 + len(line) // 110 for line in str(value).split("\n"))
        ws.row_dimensions[row].height = max(15, 15 * lines)
        row += 1

    row += 1
    title = "Average score per performance type" if len(calls) > 1 else "Score per performance type"
    _set(ws, row, 1, title, font=SECTION_FONT)
    row += 1
    _header_row(ws, row, ["Performance Type", "Score", "Max Score", "Overall Observation"])
    first = row + 1
    for metric in report.get("metrics") or []:
        row += 1
        _set(ws, row, 1, metric.get("metric_name"), alignment=WRAP, border=BORDER)
        _set(ws, row, 2, metric.get("score"), alignment=CENTER, border=BORDER, number_format="0.0")
        _set(ws, row, 3, SCORE_MAX, alignment=CENTER, border=BORDER)
        _set(ws, row, 4, metric.get("observation"), alignment=WRAP, border=BORDER)
    if row >= first:
        _score_color_scale(ws, f"B{first}:B{row}")


def _scores_sheet(ws, calls: list[dict], metric_names: list[str]):
    n = len(metric_names)
    headers = ["#", "File Name", *metric_names, "Total Score", "Max Score", "Percentage", "Call Summary"]
    _header_row(ws, 1, headers)

    first_metric_col = 3
    last_metric_col = first_metric_col + n - 1
    total_col, max_col, pct_col, summary_col = last_metric_col + 1, last_metric_col + 2, last_metric_col + 3, last_metric_col + 4
    first_letter, last_letter = get_column_letter(first_metric_col), get_column_letter(last_metric_col)

    for i, call in enumerate(calls, start=1):
        row = i + 1
        scores = {m.get("metric_name"): m.get("score") for m in call.get("metrics") or []}
        _set(ws, row, 1, i, alignment=CENTER, border=BORDER)
        _set(ws, row, 2, call.get("file_name"), alignment=WRAP, border=BORDER)
        for j, name in enumerate(metric_names):
            _set(ws, row, first_metric_col + j, scores.get(name), alignment=CENTER, border=BORDER)
        _set(ws, row, total_col, f"=SUM({first_letter}{row}:{last_letter}{row})", formula=True, alignment=CENTER, border=BORDER, font=LABEL_FONT)
        _set(ws, row, max_col, n * SCORE_MAX, alignment=CENTER, border=BORDER)
        _set(ws, row, pct_col, f"=IFERROR({get_column_letter(total_col)}{row}/{get_column_letter(max_col)}{row},0)",
             formula=True, alignment=CENTER, border=BORDER, number_format="0.0%")
        _set(ws, row, summary_col, call.get("call_summary"), alignment=WRAP, border=BORDER)

    last_row = len(calls) + 1
    if len(calls) > 1:
        avg_row = last_row + 1
        _set(ws, avg_row, 2, "Average", font=LABEL_FONT, fill=AVERAGE_FILL, border=BORDER)
        _set(ws, avg_row, 1, None, fill=AVERAGE_FILL, border=BORDER)
        for col in range(first_metric_col, pct_col + 1):
            letter = get_column_letter(col)
            _set(ws, avg_row, col, f"=AVERAGE({letter}2:{letter}{last_row})", formula=True, font=LABEL_FONT, fill=AVERAGE_FILL,
                 alignment=CENTER, border=BORDER, number_format="0.0%" if col == pct_col else "0.0")
        _set(ws, avg_row, summary_col, None, fill=AVERAGE_FILL, border=BORDER)

    if n:
        _score_color_scale(ws, f"{first_letter}2:{last_letter}{last_row + (1 if len(calls) > 1 else 0)}")

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 34
    for col in range(first_metric_col, pct_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 13
    ws.column_dimensions[get_column_letter(summary_col)].width = 80
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(summary_col)}{last_row}"


def _details_sheet(ws, calls: list[dict]):
    _header_row(ws, 1, ["File Name", "Performance Type", "Score", "Max Score", "Justification"])
    row = 1
    for call in calls:
        for metric in call.get("metrics") or []:
            row += 1
            _set(ws, row, 1, call.get("file_name"), alignment=WRAP, border=BORDER)
            _set(ws, row, 2, metric.get("metric_name"), alignment=WRAP, border=BORDER)
            _set(ws, row, 3, metric.get("score"), alignment=CENTER, border=BORDER)
            _set(ws, row, 4, SCORE_MAX, alignment=CENTER, border=BORDER)
            _set(ws, row, 5, metric.get("observation"), alignment=WRAP, border=BORDER)

    if row > 1:
        _score_color_scale(ws, f"C2:C{row}")
    for letter, width in zip("ABCDE", (34, 24, 9, 11, 100)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{max(row, 1)}"


def build_excel_report(response: dict, generated_at: Optional[datetime] = None) -> bytes:
    """Returns the .xlsx file content for an /agent-performance (or /agent-performance/upload) response."""
    report = response.get("agent_performance") or response
    calls = _calls_from_report(report)
    metric_names = [m.get("metric_name") for m in (report.get("metrics") or calls[0].get("metrics") or [])]

    wb = Workbook()
    _overview_sheet(wb.active, report, response, calls, generated_at or datetime.now())
    _scores_sheet(wb.create_sheet("Scores by Recording"), calls, metric_names)
    _details_sheet(wb.create_sheet("Detailed Justification"), calls)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
